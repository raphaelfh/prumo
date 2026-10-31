"""``build_export``: matrix omission, dictionary never refused, Celery round-trip."""

from __future__ import annotations

import io
from datetime import UTC, datetime
from unittest.mock import AsyncMock, MagicMock, patch
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.models.extraction import ExtractionCardinality, ExtractionFieldType
from app.services.exports.extraction.run import (
    XLSX_MIME,
    EmptyExportError,
    ExportFile,
    ExportRequest,
    build_export,
)
from app.services.exports.extraction.workbook import ExportShape
from app.services.extraction_export_service import (
    ArticleDescriptor,
    ExportLayout,
    ExportMode,
    ExportNotes,
    ExtractionExportService,
    FieldDescriptor,
    SectionDescriptor,
)
from app.worker.celery_app import celery_app

_PATCH = "app.services.exports.extraction.run.ExtractionExportService"


def _wide_layout(*, n_articles: int, n_subcols_each: int) -> ExportLayout:
    """A layout whose matrix needs ``n_articles * n_subcols_each`` data
    columns (+2 label columns) — used to drive the guard at/over the limit.

    The single section is ``cardinality=MANY`` so each article's
    ``section_instances`` fan out into one data column per instance, which
    is exactly what ``_article_fanout_count`` (and thus the guard) counts.
    """
    sec_id = uuid4()
    field = FieldDescriptor(
        field_id=uuid4(),
        label="F",
        type=ExtractionFieldType.TEXT,
        allowed_values=(),
    )
    section = SectionDescriptor(
        entity_type_id=sec_id,
        label="Sec",
        parent_entity_type_id=None,
        fields=(field,),
        cardinality=ExtractionCardinality.MANY,
    )
    articles = tuple(
        ArticleDescriptor(
            article_id=uuid4(),
            header_label=f"A{i}",
            run_id=uuid4(),
            version_id=uuid4(),
            entries={(sec_id, None): tuple(uuid4() for _ in range(n_subcols_each))},
        )
        for i in range(n_articles)
    )
    return ExportLayout(
        project_name="P",
        template_name="T",
        template_version=1,
        sections=(section,),
        articles=articles,
        reviewers=(),
        mode=ExportMode.CONSENSUS,
        include_ai_metadata=False,
        anonymize_reviewer_names=False,
        notes=ExportNotes(
            omitted_articles_by_stage={},
            generated_at=datetime(2026, 6, 14, tzinfo=UTC),
        ),
        value_map={},
    )


def _req(shape: ExportShape) -> ExportRequest:
    return ExportRequest(
        project_id=uuid4(),
        template_id=uuid4(),
        mode=ExportMode.CONSENSUS,
        article_ids=(),
        reviewer_id=None,
        include_ai_metadata=False,
        anonymize_reviewer_names=False,
        shape=shape,
    )


def _patched_service(layout: ExportLayout) -> MagicMock:
    service = MagicMock()
    service.return_value.resolve_layout = AsyncMock(return_value=layout)
    service.format_filename = ExtractionExportService.format_filename
    return service


async def test_build_export_omits_a_matrix_over_the_limit() -> None:
    layout = _wide_layout(n_articles=500, n_subcols_each=100)
    with patch(_PATCH, _patched_service(layout)):
        file = await build_export(MagicMock(), _req(ExportShape.COMPLETE), user_id="u")
    assert isinstance(file, ExportFile)
    assert file.media_type == XLSX_MIME
    assert "Matrix" not in load_workbook(io.BytesIO(file.content)).sheetnames


async def test_build_export_keeps_a_matrix_within_the_limit() -> None:
    layout = _wide_layout(n_articles=2, n_subcols_each=3)
    with patch(_PATCH, _patched_service(layout)):
        file = await build_export(MagicMock(), _req(ExportShape.COMPLETE), user_id="u")
    assert "Matrix" in load_workbook(io.BytesIO(file.content)).sheetnames


async def test_dictionary_export_is_never_refused_for_matrix_width() -> None:
    layout = _wide_layout(n_articles=500, n_subcols_each=100)
    with patch(_PATCH, _patched_service(layout)):
        file = await build_export(MagicMock(), _req(ExportShape.DICTIONARY), user_id="u")
    assert isinstance(file, ExportFile)
    assert "Matrix" not in load_workbook(io.BytesIO(file.content)).sheetnames


async def test_build_export_refuses_an_article_shape_with_no_articles() -> None:
    layout = _wide_layout(n_articles=0, n_subcols_each=1)
    with patch(_PATCH, _patched_service(layout)), pytest.raises(EmptyExportError):
        await build_export(MagicMock(), _req(ExportShape.COMPLETE), user_id="u")


def test_empty_export_error_survives_the_json_result_backend() -> None:
    backend = celery_app.backend
    rebuilt = backend.exception_to_python(backend.prepare_exception(EmptyExportError()))
    assert isinstance(rebuilt, EmptyExportError)
    assert str(rebuilt).startswith("EMPTY_ELIGIBLE_ARTICLES: ")
    assert rebuilt.code == "EMPTY_ELIGIBLE_ARTICLES"
    assert rebuilt.status_code == 422
