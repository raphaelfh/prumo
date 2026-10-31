import io

import pandas as pd
from openpyxl import Workbook, load_workbook

from app.services.exports.extraction.sheet_spec import (
    Cell,
    MergeSpan,
    SheetSpec,
    _render_sheet_spec,
)


def _roundtrip(spec: SheetSpec) -> bytes:
    wb = Workbook()
    ws = wb.active
    _render_sheet_spec(ws, spec)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_formula_like_and_illegal_strings_read_back_intact() -> None:
    data = _roundtrip(
        SheetSpec(
            title="S",
            rows=((Cell("h1"), Cell("h2")), (Cell("=> see table 2"), Cell("a\x0cb"))),
        )
    )
    df = pd.read_excel(io.BytesIO(data), "S", keep_default_na=False, na_values=[""])
    assert df.iloc[0, 0] == "=> see table 2"
    assert df.iloc[0, 1] == "ab"
    assert load_workbook(io.BytesIO(data))["S"]["A2"].data_type == "s"


def test_formula_like_string_in_merged_anchor_is_text() -> None:  # Review Focus 5
    spec = SheetSpec(
        title="S", rows=((Cell("=x"), Cell(None)),), merges=(MergeSpan(1, 1, 1, 2),)
    )
    ws = load_workbook(io.BytesIO(_roundtrip(spec)))["S"]
    assert ws["A1"].value == "=x" and ws["A1"].data_type == "s"
