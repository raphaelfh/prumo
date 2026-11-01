"""Codebook and Options sub-builders (spec §4.4): pure, over ``ExportLayout.codebook``."""

from __future__ import annotations

import io
from uuid import UUID, uuid4

from openpyxl import load_workbook

from app.models.extraction import ExtractionCardinality, ExtractionFieldType
from app.services.exports.codebook_model import VersionSections, build_codebook_entries
from app.services.exports.extraction.codebook import build_codebook, build_options
from app.services.exports.extraction.sheet_spec import RESERVED_TITLES, SheetSpec
from app.services.exports.extraction.workbook import (
    ExportShape,
    allocate_titles,
    build_workbook,
)
from app.services.exports.extraction_snapshot_reader import AllowedValue
from app.services.exports.values import build_option_labels
from app.services.extraction_export_service import (
    ExportLayout,
    FieldDescriptor,
    SectionDescriptor,
)
from tests.unit.helpers.export_layouts import make_layout


def _f(name: str, *options: tuple[str, str], fid: UUID | None = None, **over) -> FieldDescriptor:
    return FieldDescriptor(
        field_id=fid or uuid4(),
        label=over.pop("label", name.title()),
        type=ExtractionFieldType.SELECT if options else ExtractionFieldType.TEXT,
        allowed_values=tuple(AllowedValue(c, lbl) for c, lbl in options),
        name=name,
        **over,
    )


def _s(
    name: str, *fields: FieldDescriptor, many: bool = False, order: int = 0
) -> SectionDescriptor:
    return SectionDescriptor(
        entity_type_id=uuid4(),
        label=name.title(),
        parent_entity_type_id=None,
        fields=fields,
        cardinality=ExtractionCardinality.MANY if many else ExtractionCardinality.ONE,
        sort_order=order,
        name=name,
    )


def _with_removed_section() -> ExportLayout:
    """Active: ``study`` (ONE) + ``models`` (MANY); v1 also had ``old.x``."""
    study = _s("study", _f("design"), order=0)
    models = _s("models", _f("auc"), many=True, order=1)
    old = _s("old", _f("x"), order=2)
    active = VersionSections(uuid4(), 2, (study, models), {})
    v1 = VersionSections(uuid4(), 1, (study, models, old), {})
    labels = build_option_labels(active.sections, [v1.sections])
    return make_layout(
        sections=(study, models), codebook=build_codebook_entries(active, (v1,), labels)
    )


def _column(spec: SheetSpec, name: str) -> dict[str, object]:
    header = [c.value for c in spec.rows[0]]
    return {row[0].value: row[header.index(name)].value for row in spec.rows[1:]}


def test_codebook_header_row():
    layout = make_layout(sections=(_s("study", _f("design")),))
    titles = allocate_titles(layout, ExportShape.COMPLETE, include_matrix=True)
    assert [c.value for c in build_codebook(layout, titles).rows[0]] == [
        "field_code",
        "field_id",
        "sheet",
        "section_code",
        "section_label",
        "parent_section_code",
        "field_label",
        "type",
        "unit",
        "allowed_units",
        "required",
        "is_entity_key",
        "allowed_absent_codes",
        "allow_other",
        "description",
        "in_current_version",
        "versions",
    ]


def test_codebook_row_writes_lists_joined_and_booleans_as_bool():
    fid = uuid4()
    field = _f(
        "dose",
        fid=fid,
        label="Dose",
        unit="mg",
        allowed_units=("mg", "g"),
        is_required=True,
        is_entity_key=True,
        allowed_absent_codes=("no_information", "not_applicable"),
        description="Dose given",
    )
    layout = make_layout(sections=(_s("arm", field),))
    titles = allocate_titles(layout, ExportShape.COMPLETE, include_matrix=True)
    row = [c.value for c in build_codebook(layout, titles).rows[1]]
    assert row == [
        "arm.dose",
        str(fid),
        "Articles",
        "arm",
        "Arm",
        None,
        "Dose",
        "text",
        "mg",
        "mg; g",
        True,
        True,
        "no_information; not_applicable",
        False,
        "Dose given",
        True,
        "1",
    ]


def test_codebook_sheet_is_the_review_title_even_in_dictionary_shape():
    layout = _with_removed_section()
    titles = allocate_titles(layout, ExportShape.DICTIONARY, include_matrix=False)
    sheet_col = _column(build_codebook(layout, titles), "sheet")
    assert sheet_col == {"study.design": "Articles", "models.auc": "models", "old.x": None}


def test_section_coded_like_a_reserved_title_is_suffixed_in_codebook():  # Review Focus 1
    sec = _s("options", _f("x"), many=True)
    layout = make_layout(sections=(sec,))
    t = allocate_titles(layout, ExportShape.COMPLETE, include_matrix=True)
    title = t.levels[sec.entity_type_id]
    assert title.casefold() not in {r.casefold() for r in RESERVED_TITLES}
    assert _column(build_codebook(layout, t), "sheet")["options.x"] == title


def test_codebook_versions_are_listed_ascending():
    layout = _with_removed_section()
    titles = allocate_titles(layout, ExportShape.COMPLETE, include_matrix=True)
    versions = _column(build_codebook(layout, titles), "versions")
    assert versions == {"study.design": "1; 2", "models.auc": "1; 2", "old.x": "1"}


def test_options_keep_the_none_code():
    layout = make_layout(sections=(_s("study", _f("smoking", ("None", "None reported"))),))
    titles = allocate_titles(layout, ExportShape.COMPLETE, include_matrix=True)
    spec = build_options(layout, titles)
    assert spec is not None
    assert [c.value for c in spec.rows[0]] == [
        "field_code",
        "code",
        "label",
        "order",
        "in_current_version",
    ]
    (row,) = ([c.value for c in r] for r in spec.rows[1:])
    assert row == ["study.smoking", "None", "None reported", 1, True]


def test_options_is_absent_without_any_coded_field():
    layout = make_layout(sections=(_s("study", _f("design")),))
    titles = allocate_titles(layout, ExportShape.COMPLETE, include_matrix=True)
    assert "options" not in titles.fixed
    assert build_options(layout, titles) is None


def test_dictionary_workbook_sheets():
    layout = make_layout(sections=(_s("study", _f("smoking", ("Y", "Yes"))),))
    wb = load_workbook(
        io.BytesIO(build_workbook(layout, ExportShape.DICTIONARY, include_matrix=False))
    )
    assert wb.sheetnames == ["README", "Codebook", "Options"]
    t = allocate_titles(layout, ExportShape.DICTIONARY, include_matrix=False)
    assert t.order == tuple(wb.sheetnames)
