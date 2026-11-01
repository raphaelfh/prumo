"""The README reading snippets keep the "None" / "N/A" codes intact (spec §4.5)."""

from __future__ import annotations

import io
from uuid import uuid4

import pandas as pd
from openpyxl import load_workbook

from app.models.extraction import ExtractionCardinality, ExtractionFieldType
from app.services.exports.descriptors import value_key
from app.services.exports.extraction.data_sheet import (
    DATA_COLUMNS,
    DATA_TEXT_COLUMNS,
    DATA_TYPED_COLUMNS,
    R_READ_SNIPPET,
    pandas_read_snippet,
)
from app.services.exports.extraction.workbook import ExportShape, build_workbook
from app.services.exports.extraction_snapshot_reader import AllowedValue
from app.services.extraction_export_service import FieldDescriptor, SectionDescriptor
from tests.unit.helpers.export_layouts import export_value, make_article, make_cell, make_layout

FIELD = FieldDescriptor(
    field_id=uuid4(),
    label="Blinding",
    type=ExtractionFieldType.SELECT,
    allowed_values=(AllowedValue("None", "None"), AllowedValue("N/A", "N/A")),
    name="blinding",
)
SECTION = SectionDescriptor(
    entity_type_id=uuid4(),
    label="Study",
    parent_entity_type_id=None,
    fields=(FIELD,),
    cardinality=ExtractionCardinality.ONE,
    sort_order=0,
    name="study",
)
_A = make_article(section_id=SECTION.entity_type_id, label="Smith 2020")
_B = make_article(section_id=SECTION.entity_type_id, label="Jones 2021")
ARTICLES = (_A[0], _B[0])
_KEYS = {
    "None": value_key(_A[1], _A[2], FIELD.field_id),
    "N/A": value_key(_B[1], _B[2], FIELD.field_id),
}
VALUES = {key: export_value(value_text=code, value_label=code) for code, key in _KEYS.items()}
CELLS = {key: make_cell({"value": code}, actor_id=uuid4()) for code, key in _KEYS.items()}
LABELS = {(FIELD.field_id, "None"): "None", (FIELD.field_id, "N/A"): "N/A"}


def _xlsx() -> bytes:
    layout = make_layout(
        sections=(SECTION,),
        articles=ARTICLES,
        value_map=VALUES,
        cells=CELLS,
        option_labels=LABELS,
    )
    return build_workbook(layout, ExportShape.COMPLETE, include_matrix=True)


def test_none_and_na_codes_read_back_intact_with_readme_snippet() -> None:
    xlsx = _xlsx()
    readme = {
        c.value
        for row in load_workbook(io.BytesIO(xlsx))["README"].iter_rows()
        for c in row
        if isinstance(c.value, str)
    }
    assert pandas_read_snippet() in readme
    assert R_READ_SNIPPET in readme
    df = pd.read_excel(
        io.BytesIO(xlsx),
        "Data",
        keep_default_na=False,
        na_values=[""],
        dtype=dict.fromkeys(DATA_TEXT_COLUMNS, str),
    )
    assert sorted(df["value_text"].tolist()) == ["N/A", "None"]
    assert sorted(df["value_label"].tolist()) == ["N/A", "None"]
    # pandas defaults lose both codes.
    assert pd.read_excel(io.BytesIO(xlsx), "Data")["value_text"].isna().all()


def test_options_codes_read_back_intact() -> None:
    opts = pd.read_excel(
        io.BytesIO(_xlsx()), "Options", keep_default_na=False, na_values=[""], dtype=str
    )
    assert {"None", "N/A"} <= set(opts["code"])


def test_text_columns_cover_every_data_column_but_the_typed_ones() -> None:
    assert set(DATA_TEXT_COLUMNS) | DATA_TYPED_COLUMNS == set(DATA_COLUMNS)
    assert set(DATA_COLUMNS) >= DATA_TYPED_COLUMNS
    assert not set(DATA_TEXT_COLUMNS) & DATA_TYPED_COLUMNS


def test_pandas_snippet_names_every_text_column_as_str() -> None:
    snippet = pandas_read_snippet()
    assert snippet.startswith('pd.read_excel(path, "Data", keep_default_na=False, na_values=[""], ')
    assert '"value_text": str' in snippet
    assert '"value_num"' not in snippet
