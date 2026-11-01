"""§11 row 4 on synthetic templates: exact sheet set and key integrity.

The 30-section flat template and a nested MANY under a fieldless MANY. The
seeded templates (CHARMS, CHARMS+MM, PROBAST+AI) run the same assertions in
``tests/integration/test_extraction_export_level_fixtures.py``.
"""

from __future__ import annotations

import io
from uuid import UUID, uuid4

from openpyxl import Workbook, load_workbook

from app.models.extraction import ExtractionCardinality, ExtractionFieldType
from app.services.exports.extraction.sheet_spec import RESERVED_TITLES
from app.services.exports.extraction.workbook import ExportShape, build_workbook, matrix_fits
from app.services.extraction_export_service import (
    ArticleDescriptor,
    ExportLayout,
    FieldDescriptor,
    SectionDescriptor,
)
from tests.unit.helpers.export_layouts import make_layout


def _f(name: str) -> FieldDescriptor:
    return FieldDescriptor(
        field_id=uuid4(),
        label=name.title(),
        type=ExtractionFieldType.TEXT,
        allowed_values=(),
        name=name,
    )


def _s(
    name: str,
    *fields: FieldDescriptor,
    parent: SectionDescriptor | None = None,
    many: bool = False,
    order: int = 0,
) -> SectionDescriptor:
    return SectionDescriptor(
        entity_type_id=uuid4(),
        label=name.title(),
        parent_entity_type_id=parent.entity_type_id if parent else None,
        fields=fields,
        cardinality=ExtractionCardinality.MANY if many else ExtractionCardinality.ONE,
        sort_order=order,
        name=name,
    )


def _workbook(layout: ExportLayout) -> Workbook:
    data = build_workbook(layout, ExportShape.COMPLETE, include_matrix=matrix_fits(layout))
    return load_workbook(io.BytesIO(data))


def assert_unique_records_and_resolving_parents(wb: Workbook, article_ids: set[UUID]) -> None:
    """Every level sheet's ``record_id`` is unique and every parent resolves.

    A parent is an article or a record on some level sheet; every Codebook
    ``sheet`` names a sheet in the workbook.
    """
    levels = [n for n in wb.sheetnames if n not in RESERVED_TITLES]
    ids = {str(a) for a in article_ids}
    for name in levels:
        hdr, *rows = list(wb[name].values)
        rid = [r[hdr.index("record_id")] for r in rows]
        assert rows, name
        assert len(rid) == len(set(rid)), name
        ids |= set(rid)
    for name in levels:
        hdr, *rows = list(wb[name].values)
        assert {r[hdr.index("parent_record_id")] for r in rows} <= ids, name
    hdr, *rows = list(wb["Codebook"].values)
    assert rows
    assert {r[hdr.index("sheet")] for r in rows} - {None} <= set(wb.sheetnames)


def test_thirty_flat_root_sections_give_articles_only() -> None:
    sections = tuple(_s(f"s{i:02d}", _f(f"f{i:02d}"), order=i) for i in range(1, 31))
    article = ArticleDescriptor(
        article_id=uuid4(),
        article_label="Smith 2020",
        run_id=uuid4(),
        version_id=None,
        entries={(s.entity_type_id, None): (uuid4(),) for s in sections},
    )
    wb = _workbook(make_layout(sections=sections, articles=(article,)))
    assert wb.sheetnames == ["README", "Summary", "Matrix", "Articles", "Codebook"]
    hdr = next(wb["Articles"].values)
    assert [h for h in hdr if h.startswith("s")] == [f"s{i:02d}.f{i:02d}" for i in range(1, 31)]
    assert_unique_records_and_resolving_parents(wb, {article.article_id})


def test_nested_many_under_a_fieldless_many_gives_both_levels() -> None:
    groups = _s("groups", many=True)
    items = _s("items", _f("dose"), parent=groups, many=True)
    g1, g2 = uuid4(), uuid4()
    article = ArticleDescriptor(
        article_id=uuid4(),
        article_label="Smith 2020",
        run_id=uuid4(),
        version_id=None,
        entries={
            (groups.entity_type_id, None): (g1, g2),
            (items.entity_type_id, g1): (uuid4(), uuid4()),
            (items.entity_type_id, g2): (uuid4(), uuid4()),
        },
    )
    wb = _workbook(make_layout(sections=(groups, items), articles=(article,)))
    assert wb.sheetnames == [
        "README",
        "Summary",
        "Matrix",
        "Articles",
        "groups",
        "items",
        "Codebook",
    ]
    assert wb["groups"].max_row == 3
    assert wb["items"].max_row == 5
    hdr, *rows = list(wb["items"].values)
    assert {r[hdr.index("parent_record_id")] for r in rows} == {str(g1), str(g2)}
    assert_unique_records_and_resolving_parents(wb, {article.article_id})
