"""Matrix include-or-omit against Excel's 16,384-column hard limit.

``matrix_fits`` decides whether the Matrix sheet can be written; an oversized
Matrix is omitted (with a README note) instead of refusing the export.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime
from uuid import uuid4

from openpyxl import load_workbook

from app.models.extraction import (
    ExtractionCardinality,
    ExtractionFieldType,
)
from app.services.exports.extraction.workbook import (
    EXCEL_MAX_COLUMNS,
    ExportShape,
    allocate_titles,
    build_workbook,
    matrix_fits,
)
from app.services.extraction_export_service import (
    ArticleDescriptor,
    ExportLayout,
    ExportMode,
    FieldDescriptor,
    SectionDescriptor,
)


def _wide_layout(*, n_articles: int, n_subcols_each: int) -> ExportLayout:
    """A layout whose matrix needs ``n_articles * n_subcols_each`` data
    columns (+2 label columns) — used to drive the guard at/over the limit.

    The single section is ``cardinality=MANY`` so each article's
    ``section_instances`` fan out into one data column per instance, which
    is exactly what ``_article_fanout_count`` (and thus the guard) counts.
    """
    sec_id = uuid4()
    field = FieldDescriptor(
        field_id=uuid4(),
        label="F",
        type=ExtractionFieldType.TEXT,
        allowed_values=(),
    )
    section = SectionDescriptor(
        entity_type_id=sec_id,
        label="Sec",
        parent_entity_type_id=None,
        fields=(field,),
        cardinality=ExtractionCardinality.MANY,
    )
    articles = tuple(
        ArticleDescriptor(
            article_id=uuid4(),
            article_label=f"A{i}",
            run_id=uuid4(),
            version_id=uuid4(),
            entries={(sec_id, None): tuple(uuid4() for _ in range(n_subcols_each))},
        )
        for i in range(n_articles)
    )
    return ExportLayout(
        project_name="P",
        template_name="T",
        template_version=1,
        sections=(section,),
        articles=articles,
        reviewers=(),
        mode=ExportMode.CONSENSUS,
        include_ai_metadata=False,
        anonymize_reviewer_names=False,
        generated_at=datetime(2026, 6, 14, tzinfo=UTC),
        value_map={},
    )


def test_excel_max_columns_is_the_xlsx_hard_limit() -> None:
    assert EXCEL_MAX_COLUMNS == 16_384


def test_matrix_fits_at_the_boundary_and_not_past_it() -> None:
    assert matrix_fits(_wide_layout(n_articles=1, n_subcols_each=16_382)) is True
    assert matrix_fits(_wide_layout(n_articles=1, n_subcols_each=16_383)) is False


def test_oversized_matrix_is_omitted_with_a_readme_note() -> None:
    layout = _wide_layout(n_articles=500, n_subcols_each=100)
    wb = load_workbook(
        io.BytesIO(build_workbook(layout, ExportShape.COMPLETE, include_matrix=matrix_fits(layout)))
    )
    assert "Matrix" not in wb.sheetnames
    assert any(
        "Matrix omitted" in str(c) for r in wb["README"].iter_rows(values_only=True) for c in r if c
    )


def test_dictionary_shape_never_notes_a_matrix() -> None:
    layout = _wide_layout(n_articles=500, n_subcols_each=100)
    titles = allocate_titles(layout, ExportShape.DICTIONARY, include_matrix=False)
    assert titles.matrix_omitted is False
    assert "Matrix" not in titles.order


def test_in_bounds_layout_keeps_its_matrix_and_has_no_note() -> None:
    layout = _wide_layout(n_articles=2, n_subcols_each=3)
    wb = load_workbook(
        io.BytesIO(build_workbook(layout, ExportShape.COMPLETE, include_matrix=matrix_fits(layout)))
    )
    assert "Matrix" in wb.sheetnames
    assert not any(
        "Matrix omitted" in str(c) for r in wb["README"].iter_rows(values_only=True) for c in r if c
    )
