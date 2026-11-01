"""Pure unit tests for _build_appraisal_model (§7) — no DB."""

from __future__ import annotations

import io
import uuid
from dataclasses import replace

from openpyxl import load_workbook

from app.models.extraction import (
    ExtractionCardinality,
    ExtractionFieldType,
)
from app.services.exports.descriptors import value_key
from app.services.exports.extraction.workbook import ExportShape, build_workbook
from app.services.exports.extraction_snapshot_reader import AllowedValue
from app.services.exports.values import (
    ExportValue,
    build_option_labels,
    display_value,
    to_export_value,
)
from app.services.extraction_export_service import (
    ArticleDescriptor,
    ExportMode,
    ExtractionExportService,
    FieldDescriptor,
    ReviewerDescriptor,
    SectionDescriptor,
)
from tests.unit.helpers.export_layouts import make_layout


def _pairs(*labels: str) -> tuple[AllowedValue, ...]:
    return tuple(AllowedValue(x, x) for x in labels)


def _field(
    label,
    ftype=ExtractionFieldType.SELECT,
    allowed_values=_pairs("Low", "Unclear", "High"),
):
    return FieldDescriptor(
        field_id=uuid.uuid4(),
        label=label,
        type=ftype,
        allowed_values=allowed_values,
    )


def _ev(field: FieldDescriptor, code: str) -> ExportValue:
    return to_export_value({"value": code}, field, {})


_MARKER = {
    "No information": {"value": None, "absent_reason": "no_information"},
    "Not applicable": {"value": None, "absent_reason": "not_applicable"},
}


def _shown(model, row) -> tuple:
    return tuple(
        None if v is None else display_value(v, f)
        for v, f in zip(row.domain_verdicts, model.domain_fields, strict=True)
    )


def _section(label, verdict_field, sort_order):
    sid = uuid.uuid4()
    return SectionDescriptor(
        entity_type_id=sid,
        label=label,
        parent_entity_type_id=None,
        fields=(verdict_field,),
        cardinality=ExtractionCardinality.ONE,
        sort_order=sort_order,
    )


def test_build_appraisal_model_consensus_rollup() -> None:
    d1 = _section("Participants", _field("RoB"), 0)
    d2 = _section("Predictors", _field("RoB"), 1)
    f1 = d1.fields[0].field_id
    f2 = d2.fields[0].field_id

    run_id = uuid.uuid4()
    inst1 = uuid.uuid4()
    inst2 = uuid.uuid4()
    aid = uuid.uuid4()
    article = ArticleDescriptor(
        article_id=aid,
        article_label="Gaca, 2011",
        run_id=run_id,
        version_id=None,
        entries={(d1.entity_type_id, None): (inst1,), (d2.entity_type_id, None): (inst2,)},
    )
    value_map = {
        value_key(run_id, inst1, f1): _ev(d1.fields[0], "Low"),
        value_key(run_id, inst2, f2): _ev(d2.fields[0], "High"),
    }

    model = ExtractionExportService._build_appraisal_model(
        sections=(d1, d2),
        articles=(article,),
        reviewers=(),
        value_map=value_map,
        mode=ExportMode.CONSENSUS,
        reviewer_id=None,
    )
    assert model is not None
    assert model.domain_labels == ("Participants", "Predictors")
    assert len(model.rows) == 1
    row = model.rows[0]
    assert row.record_label == "Gaca, 2011"
    assert _shown(model, row) == ("Low", "High")
    assert row.overall == "High"  # worst-case
    assert row.per_reviewer_overall == {}


def test_build_appraisal_model_excludes_disposition_marker_verdict() -> None:
    # ADR-0016 Phase 4: a coded-disposition verdict renders (display_value)
    # as its stable label — "No information" here. The
    # domain CELL keeps that label, but the worst-case Overall EXCLUDES it, so the
    # real "High" wins instead of a fabricated Critical/most-severe. An all-marker
    # record rolls up to a blank Overall while its cells still show the label.
    d1 = _section("Participants", _field("RoB"), 0)
    d2 = _section("Predictors", _field("RoB"), 1)
    f1 = d1.fields[0].field_id
    f2 = d2.fields[0].field_id

    run_id = uuid.uuid4()
    inst1, inst2, aid = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    article = ArticleDescriptor(
        article_id=aid,
        article_label="Gaca, 2011",
        run_id=run_id,
        version_id=None,
        entries={(d1.entity_type_id, None): (inst1,), (d2.entity_type_id, None): (inst2,)},
    )

    # Mixed: one domain silent (marker → "No information"), one real "High".
    mixed = ExtractionExportService._build_appraisal_model(
        sections=(d1, d2),
        articles=(article,),
        reviewers=(),
        value_map={
            value_key(run_id, inst1, f1): to_export_value(
                _MARKER["No information"], d1.fields[0], {}
            ),
            value_key(run_id, inst2, f2): _ev(d2.fields[0], "High"),
        },
        mode=ExportMode.CONSENSUS,
        reviewer_id=None,
    )
    assert mixed is not None
    row = mixed.rows[0]
    assert _shown(mixed, row) == ("No information", "High")  # cells stay populated
    assert row.overall == "High"  # marker excluded — real verdict wins

    # All-marker: every domain silent → blank Overall, cells still labelled.
    all_marker = ExtractionExportService._build_appraisal_model(
        sections=(d1, d2),
        articles=(article,),
        reviewers=(),
        value_map={
            value_key(run_id, inst1, f1): to_export_value(
                _MARKER["No information"], d1.fields[0], {}
            ),
            value_key(run_id, inst2, f2): to_export_value(
                _MARKER["Not applicable"], d2.fields[0], {}
            ),
        },
        mode=ExportMode.CONSENSUS,
        reviewer_id=None,
    )
    assert all_marker is not None
    am_row = all_marker.rows[0]
    assert _shown(all_marker, am_row) == ("No information", "Not applicable")
    assert am_row.overall is None  # nothing assessable => blank, not most-severe


def test_build_appraisal_model_all_users_per_reviewer() -> None:
    d1 = _section("Participants", _field("RoB"), 0)
    f1 = d1.fields[0].field_id
    run_id, inst1, aid = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    r1, r2 = uuid.uuid4(), uuid.uuid4()
    reviewers = (
        ReviewerDescriptor(reviewer_id=r1, display_label="R1"),
        ReviewerDescriptor(reviewer_id=r2, display_label="R2"),
    )
    article = ArticleDescriptor(
        article_id=aid,
        article_label="Gaca, 2011",
        run_id=run_id,
        version_id=None,
        entries={(d1.entity_type_id, None): (inst1,)},
    )
    # all_users: the consensus slot is rater None; each reviewer has its own.
    value_map = {
        value_key(run_id, inst1, f1): _ev(d1.fields[0], "Low"),
        value_key(run_id, inst1, f1, r1): _ev(d1.fields[0], "Low"),
        value_key(run_id, inst1, f1, r2): _ev(d1.fields[0], "High"),
    }
    model = ExtractionExportService._build_appraisal_model(
        sections=(d1,),
        articles=(article,),
        reviewers=reviewers,
        value_map=value_map,
        mode=ExportMode.ALL_USERS,
        reviewer_id=None,
    )
    row = model.rows[0]
    assert row.overall == "Low"  # consensus rollup
    assert row.per_reviewer_overall == {r1: "Low", r2: "High"}


def test_build_appraisal_model_none_when_no_select_field() -> None:
    # A QA template whose domains carry no SELECT verdict field -> None (no sheet).
    text_field = _field("Notes", ftype=ExtractionFieldType.TEXT)
    d1 = _section("Free notes", text_field, 0)
    model = ExtractionExportService._build_appraisal_model(
        sections=(d1,),
        articles=(),
        reviewers=(),
        value_map={},
        mode=ExportMode.CONSENSUS,
        reviewer_id=None,
    )
    assert model is None


def test_build_appraisal_model_skips_signalling_select_picks_risk_label_field() -> None:
    # Mirrors the real seed: SELECT-typed signalling questions precede the
    # risk_of_bias judgment in sort_order. The verdict must be the risk-label
    # SELECT (Low/High/Unclear), NOT the first SELECT (a signalling answer).
    sid = uuid.uuid4()
    signalling = FieldDescriptor(
        field_id=uuid.uuid4(),
        label="q1_1 appropriate data sources",
        type=ExtractionFieldType.SELECT,
        allowed_values=_pairs("Y", "PY", "PN", "N", "NI", "NA"),  # _PROBAST_SIGNALING
    )
    risk = FieldDescriptor(
        field_id=uuid.uuid4(),
        label="Risk of bias",
        type=ExtractionFieldType.SELECT,
        allowed_values=_pairs("Low", "High", "Unclear"),  # _PROBAST_JUDGMENT
    )
    applicability = FieldDescriptor(
        field_id=uuid.uuid4(),
        label="Applicability concerns",
        type=ExtractionFieldType.SELECT,
        allowed_values=_pairs("Low", "High", "Unclear"),
    )
    d1 = SectionDescriptor(
        entity_type_id=sid,
        label="Participants",
        parent_entity_type_id=None,
        fields=(signalling, risk, applicability),  # signalling first, by sort_order
        cardinality=ExtractionCardinality.ONE,
        sort_order=0,
    )
    run_id, inst, aid = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    article = ArticleDescriptor(
        article_id=aid,
        article_label="Gaca, 2011",
        run_id=run_id,
        version_id=None,
        entries={(sid, None): (inst,)},
    )
    # If selection wrongly picked the signalling field, "Y" would be read and
    # ranked maximally severe -> Overall "Y". Keying on the risk-label set
    # reads risk_of_bias instead, so the verdict is the judgment "Low".
    value_map = {
        value_key(run_id, inst, signalling.field_id): _ev(signalling, "Y"),
        value_key(run_id, inst, risk.field_id): _ev(risk, "Low"),
        value_key(run_id, inst, applicability.field_id): _ev(applicability, "High"),
    }
    model = ExtractionExportService._build_appraisal_model(
        sections=(d1,),
        articles=(article,),
        reviewers=(),
        value_map=value_map,
        mode=ExportMode.CONSENSUS,
        reviewer_id=None,
    )
    assert model is not None
    row = model.rows[0]
    assert _shown(model, row) == ("Low",)  # risk_of_bias, not "Y" or "High"
    assert row.overall == "Low"


def test_v2_shapes_contribute_columns_only_for_the_judgment_section() -> None:
    """PROBAST+AI v2 (spec 2026-08-22 §8): the three D4 type sections carry
    only Y/PY/PN/N selects, assessment_scope a non-risk select, and
    overall_judgement text boxes — none is a verdict column. The single
    eval-D4 judgment section contributes exactly one."""
    run_id = uuid.uuid4()

    def sec(name, label, fields, order):
        return SectionDescriptor(
            entity_type_id=uuid.uuid4(),
            label=label,
            parent_entity_type_id=None,
            fields=tuple(fields),
            name=name,
            sort_order=order,
        )

    def fld(name, ftype, allowed):
        return FieldDescriptor(
            field_id=uuid.uuid4(),
            label=name,
            type=ftype,
            allowed_values=allowed,
            name=name,
        )

    d4_type = sec(
        "eval_d4_analysis_internal",
        "Evaluation D4 (internal)",
        [
            fld(
                "q2_reasonable_sample_size",
                ExtractionFieldType.SELECT,
                _pairs("Y", "PY", "PN", "N"),
            )
        ],
        1,
    )
    scope = sec(
        "assessment_scope",
        "Assessment scope",
        [
            fld(
                "study_type",
                ExtractionFieldType.SELECT,
                _pairs("development_only", "evaluation_only", "combination"),
            )
        ],
        2,
    )
    judgment = sec(
        "eval_d4_judgment",
        "Evaluation D4 — judgment",
        [fld("risk_of_bias", ExtractionFieldType.SELECT, _pairs("Low", "High", "Unclear"))],
        3,
    )
    overall = sec(
        "overall_judgement",
        "Overall judgement",
        [fld("summary_rob_evaluation", ExtractionFieldType.TEXT, ())],
        4,
    )
    article = ArticleDescriptor(
        article_id=uuid.uuid4(),
        article_label="A",
        run_id=run_id,
        version_id=None,
        entries={(judgment.entity_type_id, None): (uuid.uuid4(),)},
    )
    model = ExtractionExportService._build_appraisal_model(
        sections=(d4_type, scope, judgment, overall),
        articles=(article,),
        reviewers=(),
        value_map={},
        mode=ExportMode.CONSENSUS,
        reviewer_id=None,
    )
    assert model is not None
    assert model.domain_labels == ("Evaluation D4 — judgment",)


def test_coded_verdict_options_are_selected_by_their_labels() -> None:
    # The risk vocabulary lives in the labels; the stored codes may be anything.
    verdict = _field(
        "Risk of bias",
        allowed_values=(
            AllowedValue("L", "Low"),
            AllowedValue("H", "High"),
            AllowedValue("U", "Unclear"),
        ),
    )
    d1 = _section("Participants", verdict, 0)
    model = ExtractionExportService._build_appraisal_model(
        sections=(d1,),
        articles=(),
        reviewers=(),
        value_map={},
        mode=ExportMode.CONSENSUS,
        reviewer_id=None,
    )
    assert model is not None
    assert model.domain_labels == ("Participants",)


def test_overall_ranks_coded_verdicts_by_label() -> None:
    lhu = (AllowedValue("L", "Low"), AllowedValue("U", "Unclear"), AllowedValue("H", "High"))
    d1 = _section("D1", _field("RoB", allowed_values=lhu), 0)
    d2 = _section("D2", _field("RoB", allowed_values=lhu), 1)
    labels = build_option_labels([d1, d2])
    run_id, i1, i2 = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    article = ArticleDescriptor(
        article_id=uuid.uuid4(),
        article_label="A",
        run_id=run_id,
        version_id=None,
        entries={(d1.entity_type_id, None): (i1,), (d2.entity_type_id, None): (i2,)},
    )
    model = ExtractionExportService._build_appraisal_model(
        sections=(d1, d2),
        articles=(article,),
        reviewers=(),
        value_map={
            value_key(run_id, i1, d1.fields[0].field_id): to_export_value(
                {"value": "L"}, d1.fields[0], labels
            ),
            value_key(run_id, i2, d2.fields[0].field_id): to_export_value(
                {"value": "U"}, d2.fields[0], labels
            ),
        },
        mode=ExportMode.CONSENSUS,
        reviewer_id=None,
    )
    assert model is not None
    assert model.domain_fields == (d1.fields[0], d2.fields[0])
    assert tuple(v.value_text for v in model.rows[0].domain_verdicts) == ("L", "U")
    assert model.rows[0].overall == "Unclear"


def test_single_user_verdicts_read_the_reviewer_slot() -> None:
    d1 = _section("D1", _field("RoB", allowed_values=(AllowedValue("H", "High"),)), 0)
    run_id, i1, r = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    article = ArticleDescriptor(
        article_id=uuid.uuid4(),
        article_label="A",
        run_id=run_id,
        version_id=None,
        entries={(d1.entity_type_id, None): (i1,)},
    )
    model = ExtractionExportService._build_appraisal_model(
        sections=(d1,),
        articles=(article,),
        reviewers=(),
        value_map={
            value_key(run_id, i1, d1.fields[0].field_id, r): to_export_value(
                {"value": "H"}, d1.fields[0], {}
            )
        },
        mode=ExportMode.SINGLE_USER,
        reviewer_id=r,
    )
    # An unlabelled code ranks as unknown, hence most severe.
    assert model is not None and model.rows[0].overall == "H"


# ----------------------------------------------------------------------
# Rule 8: domains at the root, detected by vocabulary
# ----------------------------------------------------------------------


def test_nested_risk_verdict_is_not_a_domain_and_is_noted() -> None:
    root = _section("Participants", _field("RoB"), 0)
    nested = replace(_section("Per model", _field("RoB"), 1), parent_entity_type_id=uuid.uuid4())
    model = ExtractionExportService._build_appraisal_model(
        sections=(root, nested),
        articles=(),
        reviewers=(),
        value_map={},
        mode=ExportMode.CONSENSUS,
        reviewer_id=None,
    )
    assert model is not None
    assert model.domain_labels == ("Participants",)
    assert ExtractionExportService._appraisal_notes(
        (root, nested), template_kind="extraction", appraisal=model
    ) == ("Appraisal: 'Per model' is not at the root, so it is not an Appraisal domain.",)


def test_only_nested_risk_verdicts_give_no_appraisal() -> None:
    nested = replace(_section("Per model", _field("RoB"), 0), parent_entity_type_id=uuid.uuid4())
    model = ExtractionExportService._build_appraisal_model(
        sections=(nested,),
        articles=(),
        reviewers=(),
        value_map={},
        mode=ExportMode.CONSENSUS,
        reviewer_id=None,
    )
    assert model is None


def test_custom_vocabulary_qa_template_has_no_appraisal_and_a_note() -> None:
    d = _section("Quality", _field("Grade", allowed_values=_pairs("Good", "Fair", "Poor")), 0)
    model = ExtractionExportService._build_appraisal_model(
        sections=(d,),
        articles=(),
        reviewers=(),
        value_map={},
        mode=ExportMode.CONSENSUS,
        reviewer_id=None,
    )
    assert model is None
    notes = ExtractionExportService._appraisal_notes(
        (d,), template_kind="quality_assessment", appraisal=None
    )
    assert notes == (
        "Appraisal: no root section has a SELECT field whose labels all belong to the "
        "risk-of-bias vocabulary (PROBAST, QUADAS, RoB), so there is no Appraisal sheet.",
    )


def test_extraction_template_without_verdicts_gets_no_note() -> None:
    d = _section("Quality", _field("Grade", allowed_values=_pairs("Good", "Fair", "Poor")), 0)
    assert (
        ExtractionExportService._appraisal_notes((d,), template_kind="extraction", appraisal=None)
        == ()
    )


def test_workbook_without_appraisal_prints_the_note_in_the_readme() -> None:
    note = "Appraisal: no root section has a SELECT field (test note)."
    layout = make_layout(
        sections=(_section("Quality", _field("Grade"), 0),),
        appraisal=None,
        appraisal_notes=(note,),
    )
    wb = load_workbook(
        io.BytesIO(build_workbook(layout, ExportShape.COMPLETE, include_matrix=True))
    )
    assert "Appraisal" not in wb.sheetnames
    readme = [c for row in wb["README"].iter_rows(values_only=True) for c in row]
    assert note in readme
