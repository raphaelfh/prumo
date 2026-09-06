"""Golden geometry for a CHARMS-shaped export (trees B4 characterization).

Spec §10 asks for a golden CHARMS workbook that is byte-identical across
the tree generalization. Those bytes cannot be the goal: they are wrong
today. ``_load_descriptors`` appends EVERY child-section instance to one
flat ``ArticleDescriptor.model_instances`` tuple
(``extraction_export_service.py`` §"model_instances"), while
``hitl_session_service._backfill_child_singletons`` guarantees exactly one
instance per (entry, singleton child). Seeded CHARMS declares SIX child
sections, so one entry produces a six-long tuple that
``matrix._article_fanout_count`` reads as the column count and
``matrix._resolve_instance_id`` indexes positionally — one entry renders as
six sub-columns with its values on a diagonal.

So this file pins the geometry a reviewer needs instead of the bytes:
one sub-column per ENTRY, every one of that entry's sections readable in
that entry's column. It is the characterization the generalization has to
satisfy, and it fails against the flat tuple.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime
from uuid import UUID, uuid4

from openpyxl import load_workbook

from app.models.extraction import (
    ExtractionCardinality,
    ExtractionEntityRole,
    ExtractionFieldType,
)
from app.services.exports.extraction.workbook import build_workbook
from app.services.extraction_export_service import (
    ArticleDescriptor,
    ExportLayout,
    ExportMode,
    ExportNotes,
    FieldDescriptor,
    SectionDescriptor,
)

# Seeded CHARMS carries six sections under ``prediction_models``; the count
# is what turns the flat tuple into a six-step staircase, so keep it real.
_CHILD_LABELS = (
    "Model Development",
    "Model Performance",
    "Model Evaluation",
    "Outcome",
    "Predictors",
    "Sample Size",
)


class _Charms:
    """A CHARMS-shaped layout: one root group, six singleton children."""

    def __init__(self, *, entry_count: int) -> None:
        self.container_id = uuid4()
        self.child_ids = tuple(uuid4() for _ in _CHILD_LABELS)
        self.field_ids = tuple(uuid4() for _ in _CHILD_LABELS)
        self.run_id = uuid4()
        # entry -> that entry's six singleton child instances
        self.entry_ids = tuple(uuid4() for _ in range(entry_count))
        self.child_instances = {
            (entry, child): uuid4() for entry in self.entry_ids for child in self.child_ids
        }

    def sections(self) -> tuple[SectionDescriptor, ...]:
        container = SectionDescriptor(
            entity_type_id=self.container_id,
            label="Prediction Models",
            role=ExtractionEntityRole.MODEL_CONTAINER,
            parent_entity_type_id=None,
            fields=(),
            cardinality=ExtractionCardinality.MANY,
            sort_order=0,
            entry_label="model",
        )
        children = tuple(
            SectionDescriptor(
                entity_type_id=cid,
                label=label,
                role=ExtractionEntityRole.MODEL_SECTION,
                parent_entity_type_id=self.container_id,
                fields=(
                    FieldDescriptor(
                        field_id=fid,
                        label=f"{label} field",
                        type=ExtractionFieldType.TEXT,
                        allowed_values=(),
                    ),
                ),
                # Production child sections are cardinality ONE; the N-entry
                # fan-out is meant to come from the entries, not from here.
                cardinality=ExtractionCardinality.ONE,
                sort_order=i + 1,
            )
            for i, (cid, fid, label) in enumerate(
                zip(self.child_ids, self.field_ids, _CHILD_LABELS, strict=True)
            )
        )
        return (container, *children)

    def article(self) -> ArticleDescriptor:
        flat: list[UUID] = [
            self.child_instances[(entry, child)]
            for entry in self.entry_ids
            for child in self.child_ids
        ]
        entries: dict[tuple[UUID, UUID | None], tuple[UUID, ...]] = {
            (self.container_id, None): self.entry_ids
        }
        for entry in self.entry_ids:
            for child in self.child_ids:
                entries[(child, entry)] = (self.child_instances[(entry, child)],)
        return ArticleDescriptor(
            article_id=uuid4(),
            header_label="Gaca, 2011",
            run_id=self.run_id,
            version_id=None,
            # The flat projection is kept beside the tree only until its last
            # reader goes; `flat` is exactly what produced the diagonal.
            model_instances=tuple(flat),
            entries=entries,
        )

    def value(self, entry_index: int, child_index: int) -> str:
        return f"e{entry_index}-c{child_index}"

    def layout(self) -> ExportLayout:
        value_map = {
            (
                self.run_id,
                self.child_instances[(entry, self.child_ids[c])],
                self.field_ids[c],
            ): self.value(e, c)
            for e, entry in enumerate(self.entry_ids)
            for c in range(len(self.child_ids))
        }
        return ExportLayout(
            project_name="P",
            template_name="CHARMS",
            template_version=1,
            sections=self.sections(),
            articles=(self.article(),),
            reviewers=(),
            mode=ExportMode.CONSENSUS,
            include_ai_metadata=False,
            anonymize_reviewer_names=False,
            notes=ExportNotes(generated_at=datetime(2026, 6, 14, tzinfo=UTC)),
            value_map=value_map,
        )


def _sheet(charms: _Charms):
    return load_workbook(io.BytesIO(build_workbook(charms.layout())))["CHARMS"]


def _value_columns(ws) -> int:
    """Sub-columns after the two label columns (Section, Field)."""
    return ws.max_column - 2


def test_one_entry_is_one_subcolumn_with_every_section_readable():
    charms = _Charms(entry_count=1)
    ws = _sheet(charms)

    assert _value_columns(ws) == 1, "one model must occupy exactly one column"

    # Every one of the entry's six sections must be readable in that column —
    # the whole point of the sheet is reading one model down a single column.
    column = [ws.cell(row=r, column=3).value for r in range(1, ws.max_row + 1)]
    for c in range(len(_CHILD_LABELS)):
        assert charms.value(0, c) in column, (
            f"section {_CHILD_LABELS[c]!r} is not readable in the entry's column"
        )


def test_two_entries_are_two_subcolumns_each_internally_complete():
    charms = _Charms(entry_count=2)
    ws = _sheet(charms)

    assert _value_columns(ws) == 2, "two models must occupy exactly two columns"

    for e in range(2):
        column = [ws.cell(row=r, column=3 + e).value for r in range(1, ws.max_row + 1)]
        for c in range(len(_CHILD_LABELS)):
            assert charms.value(e, c) in column, (
                f"entry {e} section {_CHILD_LABELS[c]!r} missing from its own column"
            )
        # And no bleed: entry 0's column must not carry entry 1's values.
        other = 1 - e
        for c in range(len(_CHILD_LABELS)):
            assert charms.value(other, c) not in column, (
                f"entry {e}'s column carries entry {other}'s value"
            )
