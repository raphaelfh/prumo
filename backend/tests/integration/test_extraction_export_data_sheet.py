"""Integration: the Data sheet rendered from a real resolved layout (spec §4.1, §7).

Real Postgres, SEED graph. Each test builds its own run on the seed article;
the SAVEPOINT rollback drops it.
"""

from __future__ import annotations

import io
import json
from collections import Counter
from datetime import datetime
from typing import Any
from uuid import UUID, uuid4

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.extraction import ExtractionRun, ExtractionRunStage
from app.models.extraction_versioning import ExtractionTemplateVersion
from app.services.exports.extraction.workbook import ExportShape, build_workbook
from app.services.extraction_consensus_service import ExtractionConsensusService
from app.services.extraction_export_service import ExportLayout, ExportMode
from app.services.extraction_review_service import ExtractionReviewService
from app.services.extraction_snapshot import build_template_version_snapshot
from app.services.run_lifecycle_service import RunLifecycleService
from tests.integration.conftest import SEED, land_ai_proposal
from tests.integration.helpers.export_fixtures import export_service, make_run

pytestmark = [pytest.mark.asyncio, pytest.mark.integration]

_INST = SEED.primary_instance
_FIELD = SEED.primary_field


async def _extract_run(db: AsyncSession) -> ExtractionRun:
    run = await make_run(
        db,
        project_id=SEED.primary_project,
        article_id=SEED.primary_article,
        template_id=SEED.primary_template,
        profile_id=SEED.primary_profile,
    )
    if run.stage == ExtractionRunStage.PENDING.value:
        await RunLifecycleService(db).advance_stage(
            run_id=run.id, target_stage=ExtractionRunStage.EXTRACT, user_id=SEED.primary_profile
        )
    return run


async def _decide(
    db: AsyncSession, run_id: UUID, reviewer: UUID, decision: str, *, field_id: UUID = _FIELD,
    **kw: Any,
):  # fmt: skip
    return await ExtractionReviewService(db).record_decision(
        run_id=run_id, instance_id=_INST, field_id=field_id, reviewer_id=reviewer,
        decision=decision, **kw,
    )  # fmt: skip


async def _two_reviewers_in_consensus(db: AsyncSession) -> ExtractionRun:
    run = await _extract_run(db)
    await _decide(db, run.id, SEED.primary_profile, "edit", value={"value": 10})
    await _decide(db, run.id, SEED.reviewer_profile, "edit", value={"value": 12})
    await RunLifecycleService(db).advance_stage(
        run_id=run.id, target_stage=ExtractionRunStage.CONSENSUS, user_id=SEED.primary_profile
    )
    await ExtractionConsensusService(db).record_consensus(
        run_id=run.id, instance_id=_INST, field_id=_FIELD,
        consensus_user_id=SEED.primary_profile, mode="manual_override", value={"value": 11},
        rationale="split the difference",
    )  # fmt: skip
    return run


async def _layout(db: AsyncSession, mode: ExportMode, **kw: Any) -> ExportLayout:
    return await export_service(db, profile_id=SEED.primary_profile).resolve_layout(
        project_id=SEED.primary_project,
        template_id=SEED.primary_template,
        mode=mode,
        article_ids=[SEED.primary_article],
        include_ai_metadata=False,
        anonymize_reviewer_names=False,
        **kw,
    )


def _workbook(layout: ExportLayout):
    return load_workbook(
        io.BytesIO(build_workbook(layout, ExportShape.COMPLETE, include_matrix=True))
    )


def _header_col(ws: Worksheet, name: str) -> int:
    return next(c.column for c in ws[1] if c.value == name)


def _data(layout: ExportLayout) -> list[dict[str, Any]]:
    ws = _workbook(layout)["Data"]
    header = [c.value for c in ws[1]]
    return [
        dict(zip(header, (c.value for c in row), strict=True)) for row in ws.iter_rows(min_row=2)
    ]


async def test_legacy_accept_proposal_reads_answered_ai_accepted(db_session: AsyncSession) -> None:
    run = await _extract_run(db_session)
    proposal = await land_ai_proposal(
        db_session, run_id=run.id, instance_id=_INST, field_id=_FIELD, proposed_value={"value": 42}
    )
    await _decide(
        db_session, run.id, SEED.primary_profile, "accept_proposal", proposal_record_id=proposal.id
    )

    layout = await _layout(db_session, ExportMode.SINGLE_USER, reviewer_id=SEED.primary_profile)
    (row,) = [r for r in _data(layout) if r["field_id"] == str(_FIELD)]

    assert (row["status"], row["origin"], row["value_num"]) == ("answered", "ai_accepted", 42)
    assert (row["ai_proposed"], row["ai_proposal_id"]) == (True, str(proposal.id))


async def test_all_users_has_one_consensus_row_per_coordinate(db_session: AsyncSession) -> None:
    await _two_reviewers_in_consensus(db_session)

    data = _data(await _layout(db_session, ExportMode.ALL_USERS))
    coords = Counter((r["instance_id"], r["field_id"]) for r in data if r["rater"] == "consensus")
    reviewer_rows = [r for r in data if r["rater"] != "consensus"]

    assert coords and coords.most_common(1)[0][1] == 1
    assert len(reviewer_rows) == 2
    (consensus,) = [r for r in data if r["rater"] == "consensus"]
    assert (consensus["consensus_method"], consensus["origin"]) == ("override", "human")
    assert (consensus["n_raters_with_value"], consensus["rationale"]) == (2, "split the difference")


async def test_two_same_name_reviewers_have_distinct_rater(db_session: AsyncSession) -> None:
    for pid in (SEED.primary_profile, SEED.reviewer_profile):
        await db_session.execute(
            text("UPDATE public.profiles SET full_name = 'Alex Doe' WHERE id = :id"),
            {"id": str(pid)},
        )
    await _two_reviewers_in_consensus(db_session)

    data = _data(await _layout(db_session, ExportMode.ALL_USERS))
    reviewer_rows = [r for r in data if r["rater"] != "consensus"]

    assert len(reviewer_rows) == 2
    assert len({r["rater"] for r in reviewer_rows}) == 2


async def test_decided_at_cell_is_a_datetime(db_session: AsyncSession) -> None:
    await _two_reviewers_in_consensus(db_session)

    ws = _workbook(await _layout(db_session, ExportMode.ALL_USERS))["Data"]

    assert ws.max_row >= 2
    assert isinstance(ws.cell(2, _header_col(ws, "decided_at")).value, datetime)


async def test_every_data_record_id_joins_its_level_sheet(db_session: AsyncSession) -> None:
    await _two_reviewers_in_consensus(db_session)

    wb = _workbook(await _layout(db_session, ExportMode.ALL_USERS))
    ws = wb["Data"]
    header = [c.value for c in ws[1]]
    data = [dict(zip(header, (c.value for c in r), strict=True)) for r in ws.iter_rows(min_row=2)]

    joined = [r for r in data if r["sheet"]]
    assert joined  # precondition: the seed value lands on a level sheet
    for r in joined:
        level = wb[r["sheet"]]
        key = "article_id" if r["sheet"] == "Articles" else "record_id"
        col = _header_col(level, key)
        assert r["record_id"] in {row[col - 1].value for row in level.iter_rows(min_row=2)}


async def test_removed_field_value_is_present_in_data(db_session: AsyncSession) -> None:
    """A field only the run's pinned version holds: in Data, with a blank sheet."""
    snapshot = await build_template_version_snapshot(db_session, SEED.primary_template)
    await db_session.execute(
        text(
            "UPDATE public.extraction_template_versions SET schema = CAST(:s AS jsonb) "
            "WHERE project_template_id = :tid AND is_active"
        ),
        {"s": json.dumps(snapshot), "tid": str(SEED.primary_template)},
    )
    new_fid = uuid4()
    await db_session.execute(
        text(
            "INSERT INTO public.extraction_fields "
            "(id, entity_type_id, name, label, field_type, is_required) "
            "VALUES (:id, :etid, 'gone_field', 'Gone field', 'text', false)"
        ),
        {"id": str(new_fid), "etid": str(SEED.primary_entity_type)},
    )
    pinned = ExtractionTemplateVersion(
        project_template_id=SEED.primary_template,
        version=9001,
        is_active=False,
        published_by=SEED.primary_profile,
        schema_=await build_template_version_snapshot(db_session, SEED.primary_template),
    )
    db_session.add(pinned)
    await db_session.flush()
    run = await _extract_run(db_session)
    await db_session.execute(
        text("UPDATE public.extraction_runs SET version_id = :v WHERE id = :r"),
        {"v": str(pinned.id), "r": str(run.id)},
    )
    await _decide(
        db_session, run.id, SEED.primary_profile, "edit",
        field_id=new_fid, value={"value": "kept"}, enforce_contract=False,
    )  # fmt: skip

    layout = await _layout(db_session, ExportMode.SINGLE_USER, reviewer_id=SEED.primary_profile)

    assert new_fid not in {f.field_id for s in layout.sections for f in s.fields}
    (r,) = [r for r in _data(layout) if r["field_id"] == str(new_fid)]
    assert r["field_code"] and r["sheet"] in (None, "") and r["value_text"] == "kept"
