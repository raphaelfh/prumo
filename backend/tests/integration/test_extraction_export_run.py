"""``authorize_export``: validation, membership, then ownership of articles and template."""

from __future__ import annotations

import io
from typing import Any
from uuid import UUID, uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.error_handler import AuthorizationError, NotFoundError, ValidationError
from app.services.exports.extraction.run import (
    EmptyExportError,
    ExportRequest,
    authorize_export,
    build_export,
)
from app.services.exports.extraction.workbook import ExportShape
from app.services.extraction_export_service import ExportMode
from tests.integration.conftest import SEED

pytestmark = [pytest.mark.asyncio, pytest.mark.integration]


async def _insert_article(db: AsyncSession, project_id: UUID, *, title: str) -> UUID:
    article_id = uuid4()
    await db.execute(
        text(
            "INSERT INTO public.articles (id, project_id, title, row_version) "
            "VALUES (:id, :pid, :title, 1)"
        ),
        {"id": str(article_id), "pid": str(project_id), "title": title},
    )
    return article_id


def _req(**kw) -> ExportRequest:
    base: dict[str, Any] = {
        "project_id": SEED.primary_project,
        "template_id": SEED.primary_template,
        "mode": ExportMode.ALL_USERS,
        "article_ids": (SEED.primary_article,),
        "reviewer_id": None,
        "include_ai_metadata": False,
        "anonymize_reviewer_names": False,
        "shape": ExportShape.COMPLETE,
    }
    return ExportRequest(**{**base, **kw})


async def test_foreign_article_is_not_found(db_session: AsyncSession) -> None:
    foreign = await _insert_article(
        db_session, SEED.secondary_project, title="FOREIGN-SENTINEL-TITLE"
    )
    with pytest.raises(NotFoundError) as exc:
        await authorize_export(
            db_session, _req(article_ids=(foreign,)), user_id=str(SEED.primary_profile)
        )
    assert "FOREIGN-SENTINEL-TITLE" not in exc.value.message


async def test_foreign_template_is_not_found(db_session: AsyncSession) -> None:
    with pytest.raises(NotFoundError):
        await authorize_export(
            db_session,
            _req(project_id=SEED.secondary_project, article_ids=(), shape=ExportShape.DICTIONARY),
            user_id=str(SEED.primary_profile),
        )


async def test_single_user_requires_reviewer_id(db_session: AsyncSession) -> None:
    with pytest.raises(ValidationError):
        await authorize_export(
            db_session, _req(mode=ExportMode.SINGLE_USER), user_id=str(SEED.primary_profile)
        )


async def test_authorize_checks_membership_before_ownership(db_session: AsyncSession) -> None:
    foreign = await _insert_article(db_session, SEED.secondary_project, title="X")
    with pytest.raises(AuthorizationError):
        await authorize_export(
            db_session,
            _req(mode=ExportMode.CONSENSUS, article_ids=(foreign,)),
            user_id=str(SEED.outsider_profile),
        )


async def test_owner_passes(db_session: AsyncSession) -> None:
    await authorize_export(db_session, _req(), user_id=str(SEED.primary_profile))


async def test_build_export_refuses_an_empty_complete_export(db_session: AsyncSession) -> None:
    await db_session.execute(
        text("DELETE FROM public.extraction_runs WHERE project_id = :p AND article_id = :a"),
        {"p": str(SEED.primary_project), "a": str(SEED.primary_article)},
    )
    with pytest.raises(EmptyExportError) as exc:
        await build_export(
            db_session, _req(mode=ExportMode.CONSENSUS), user_id=str(SEED.primary_profile)
        )
    assert exc.value.code == "EMPTY_ELIGIBLE_ARTICLES"
    assert exc.value.message.startswith("EMPTY_ELIGIBLE_ARTICLES: ")


async def test_dictionary_export_with_no_articles_builds(db_session: AsyncSession) -> None:
    file = await build_export(
        db_session,
        _req(mode=ExportMode.CONSENSUS, article_ids=(), shape=ExportShape.DICTIONARY),
        user_id=str(SEED.primary_profile),
    )
    assert file.filename.endswith(".xlsx") and "_consensus_" in file.filename
    assert load_workbook(io.BytesIO(file.content)).sheetnames[0] == "README"
