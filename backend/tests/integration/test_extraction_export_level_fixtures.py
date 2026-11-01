"""§11 row 4 on the seeded templates: exact sheet set and key integrity.

CHARMS, CHARMS+MM and PROBAST+AI are cloned into the primary project, given
two instances per parent for every MANY section and one for every ONE section,
and exported. The synthetic templates (30 flat sections; a nested MANY under a
fieldless MANY) run the same assertions in
``tests/unit/test_extraction_export_level_fixtures.py``.
"""

from __future__ import annotations

import io
from typing import Any
from uuid import UUID, uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.extraction import ExtractionCardinality
from app.models.extraction_versioning import ExtractionTemplateVersion, TemplateKind
from app.services.exports.descriptors import walk_sections
from app.services.exports.extraction.workbook import ExportShape, build_workbook, matrix_fits
from app.services.extraction_export_service import ExportMode, ExtractionExportService
from app.services.template_versioning import clone_template
from tests.integration.conftest import SEED
from tests.unit.test_extraction_export_level_fixtures import (
    assert_unique_records_and_resolving_parents,
)

_INSERT_INSTANCE = text(
    "INSERT INTO public.extraction_instances "
    "(id, project_id, template_id, entity_type_id, article_id, parent_instance_id, "
    " label, sort_order, created_by) "
    "VALUES (:id, :pid, :tid, :etid, :aid, :parent, :label, :so, :created_by)"
)


async def _materialize(
    db: AsyncSession, svc: ExtractionExportService, clone: Any, *, kind: str
) -> None:
    """Instances for every walked section and one finalized run on the primary article.

    ``clone`` is the ``TemplateClone`` from ``clone_template``; the type lives in a
    private module this test may not import (``test_template_versioning_boundary``).
    """
    scope = {
        "pid": str(SEED.primary_project),
        "aid": str(SEED.primary_article),
        "tid": str(clone.project_template_id),
    }
    for table in ("extraction_runs", "extraction_instances"):
        await db.execute(
            text(
                f"DELETE FROM public.{table} WHERE project_id = :pid "  # noqa: S608 - fixed names
                "AND article_id = :aid AND template_id = :tid"
            ),
            scope,
        )
    version = await db.get(ExtractionTemplateVersion, clone.version_id)
    assert version is not None
    (active,) = await svc._load_pinned_versions(
        template_id=clone.project_template_id, versions=(version,)
    )
    by_section: dict[UUID, list[UUID]] = {}
    for section in walk_sections(active.sections):
        parents = (
            by_section[section.parent_entity_type_id]
            if section.parent_entity_type_id is not None
            else [None]
        )
        count = 2 if section.cardinality is ExtractionCardinality.MANY else 1
        created = by_section.setdefault(section.entity_type_id, [])
        for parent in parents:
            for i in range(count):
                instance_id = uuid4()
                created.append(instance_id)
                await db.execute(
                    _INSERT_INSTANCE,
                    {
                        **scope,
                        "id": str(instance_id),
                        "etid": str(section.entity_type_id),
                        "parent": str(parent) if parent else None,
                        "label": f"{section.name} {i + 1}",
                        "so": i,
                        "created_by": str(SEED.primary_profile),
                    },
                )
    await db.execute(
        text(
            "INSERT INTO public.extraction_runs "
            "(id, project_id, article_id, template_id, kind, version_id, "
            " stage, status, created_by) "
            "VALUES (gen_random_uuid(), :pid, :aid, :tid, :kind, :vid, "
            " 'finalized', 'completed', :created_by)"
        ),
        {
            **scope,
            "kind": kind,
            "vid": str(clone.version_id),
            "created_by": str(SEED.primary_profile),
        },
    )


@pytest.mark.parametrize(
    ("global_id", "kind", "expected"),
    [
        (
            UUID("000c0000-0000-0000-0000-000000000001"),
            TemplateKind.EXTRACTION,
            [
                "README",
                "Summary",
                "Matrix",
                "Articles",
                "prediction_models",
                "final_predictors",
                "Codebook",
                "Options",
            ],
        ),
        (
            UUID("000e0000-0000-0000-0000-000000000001"),
            TemplateKind.EXTRACTION,
            [
                "README",
                "Summary",
                "Matrix",
                "Articles",
                "prediction_models",
                "numeric_performance",
                "Codebook",
                "Options",
            ],
        ),
        (
            UUID("00ba0000-0000-0000-0000-000000000002"),
            TemplateKind.QUALITY_ASSESSMENT,
            ["README", "Summary", "Matrix", "Articles", "Appraisal", "Codebook", "Options"],
        ),
    ],
    ids=["charms", "charms_mm", "probast_ai"],
)
async def test_real_templates_give_the_exact_sheet_set_unique_records_and_resolving_parents(
    db_session: AsyncSession, global_id: UUID, kind: TemplateKind, expected: list[str]
) -> None:
    svc = ExtractionExportService(db=db_session, user_id=str(SEED.primary_profile))
    clone = await clone_template(
        db_session,
        project_id=SEED.primary_project,
        global_template_id=global_id,
        user_id=SEED.primary_profile,
        kind=kind,
    )
    await _materialize(db_session, svc, clone, kind=kind.value)
    layout = await svc.resolve_layout(
        project_id=SEED.primary_project,
        template_id=clone.project_template_id,
        mode=ExportMode.CONSENSUS,
        article_ids=[SEED.primary_article],
        include_ai_metadata=False,
        anonymize_reviewer_names=False,
    )
    data = build_workbook(layout, ExportShape.COMPLETE, include_matrix=matrix_fits(layout))
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == expected
    assert_unique_records_and_resolving_parents(wb, {SEED.primary_article})
