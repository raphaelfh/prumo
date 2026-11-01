"""Integration: an adjudicator with no reviewer decision still gets a Member letter.

Regression pin for PR5's ordering (reviewer **or consensus** decisions, spec
§3, D4): the anonymised export must letter the consensus author and never
leak a real name. Real Postgres, SEED graph; the SAVEPOINT rollback drops
every row written here.
"""

from __future__ import annotations

import io
import re
from typing import Any

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.extraction import ExtractionRunStage
from app.models.extraction_workflow import (
    ExtractionConsensusMode,
    ExtractionReviewerDecisionType,
)
from app.services.exports.extraction.run import ExportRequest, build_export
from app.services.exports.extraction.workbook import ExportShape
from app.services.extraction_consensus_service import ExtractionConsensusService
from app.services.extraction_export_service import ExportMode
from app.services.run_lifecycle_service import RunLifecycleService
from tests.integration.conftest import SEED
from tests.integration.helpers.export_fixtures import export_service, make_run, record_decision

pytestmark = [pytest.mark.asyncio, pytest.mark.integration]


def _rows_as_dicts(ws: Worksheet) -> list[dict[str, Any]]:
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    return [dict(zip(header, row, strict=True)) for row in rows[1:]]


async def test_adjudicator_without_reviewer_decisions_gets_a_member_letter(
    db_session: AsyncSession,
) -> None:
    run = await make_run(
        db_session,
        project_id=SEED.primary_project,
        article_id=SEED.primary_article,
        template_id=SEED.primary_template,
        profile_id=SEED.primary_profile,
    )
    lifecycle = RunLifecycleService(db_session)
    if run.stage == ExtractionRunStage.PENDING.value:  # PENDING may only advance to EXTRACT
        await lifecycle.advance_stage(
            run_id=run.id, target_stage=ExtractionRunStage.EXTRACT, user_id=SEED.primary_profile
        )
    await record_decision(
        db_session,
        run_id=run.id,
        instance_id=SEED.primary_instance,
        field_id=SEED.primary_field,
        reviewer_id=SEED.reviewer_profile,
        decision=ExtractionReviewerDecisionType.EDIT,
        value={"value": 10},
    )
    await lifecycle.advance_stage(
        run_id=run.id, target_stage=ExtractionRunStage.CONSENSUS, user_id=SEED.primary_profile
    )
    await ExtractionConsensusService(db_session).record_consensus(
        run_id=run.id,
        instance_id=SEED.primary_instance,
        field_id=SEED.primary_field,
        consensus_user_id=SEED.primary_profile,
        mode=ExtractionConsensusMode.MANUAL_OVERRIDE,
        value={"value": 12},
        rationale="adjudicated",
    )
    await lifecycle.advance_stage(
        run_id=run.id, target_stage=ExtractionRunStage.FINALIZED, user_id=SEED.primary_profile
    )

    labels = await export_service(db_session, profile_id=SEED.primary_profile)._load_rater_labels(
        project_id=SEED.primary_project, anonymize=True
    )
    assert re.fullmatch(r"Member [A-Z]+", labels[SEED.primary_profile])
    assert labels[SEED.primary_profile] != labels[SEED.reviewer_profile]

    out = await build_export(
        db_session,
        ExportRequest(
            project_id=SEED.primary_project,
            template_id=SEED.primary_template,
            mode=ExportMode.ALL_USERS,
            article_ids=(SEED.primary_article,),
            reviewer_id=None,
            include_ai_metadata=False,
            anonymize_reviewer_names=True,
            shape=ExportShape.COMPLETE,
        ),
        user_id=str(SEED.primary_profile),
    )
    wb = load_workbook(io.BytesIO(out.content))
    consensus = [r for r in _rows_as_dicts(wb["Data"]) if r["rater"] == "consensus"]
    assert consensus
    assert {r["decided_by"] for r in consensus} == {labels[SEED.primary_profile]}
    assert {r["consensus_by"] for r in _rows_as_dicts(wb["Verification"])} == {
        labels[SEED.primary_profile]
    }
    assert "Integration Reviewer" not in {
        c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row
    }
