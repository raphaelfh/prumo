"""Integration: NI/NA/NE abbreviations appear only in Publication tables (spec §11 row 6).

Real Postgres, SEED graph. The test builds its own finalized run on the seed
article; the SAVEPOINT rollback drops it.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.extraction import ExtractionRunStage
from app.models.extraction_workflow import ExtractionConsensusMode
from app.services.exports.extraction.run import ExportFile, ExportRequest, build_export
from app.services.exports.extraction.workbook import ExportShape
from app.services.exports.values import SCOPE_MARKER
from app.services.extraction_consensus_service import ExtractionConsensusService
from app.services.extraction_export_service import ExportMode
from app.services.run_lifecycle_service import RunLifecycleService
from tests.integration.conftest import SEED
from tests.integration.helpers.export_fixtures import make_run

pytestmark = [pytest.mark.asyncio, pytest.mark.integration]


async def _seed_finalized_absent_run(db: AsyncSession) -> None:
    """A finalized run whose consensus on the seed NUMBER field is "no information"."""
    run = await make_run(
        db,
        project_id=SEED.primary_project,
        article_id=SEED.primary_article,
        template_id=SEED.primary_template,
        profile_id=SEED.primary_profile,
    )
    lifecycle = RunLifecycleService(db)
    if run.stage == ExtractionRunStage.PENDING.value:
        await lifecycle.advance_stage(
            run_id=run.id, target_stage=ExtractionRunStage.EXTRACT, user_id=SEED.primary_profile
        )
    await lifecycle.advance_stage(
        run_id=run.id, target_stage=ExtractionRunStage.CONSENSUS, user_id=SEED.primary_profile
    )
    await ExtractionConsensusService(db).record_consensus(
        run_id=run.id,
        instance_id=SEED.primary_instance,
        field_id=SEED.primary_field,
        consensus_user_id=SEED.primary_profile,
        mode=ExtractionConsensusMode.MANUAL_OVERRIDE,
        value={"value": {"value": None, "absent_reason": "no_information"}},
        rationale="publication markers",
        enforce_contract=False,
    )
    await lifecycle.advance_stage(
        run_id=run.id, target_stage=ExtractionRunStage.FINALIZED, user_id=SEED.primary_profile
    )


def _request(shape: ExportShape) -> ExportRequest:
    return ExportRequest(
        project_id=SEED.primary_project,
        template_id=SEED.primary_template,
        mode=ExportMode.CONSENSUS,
        article_ids=(SEED.primary_article,),
        reviewer_id=None,
        include_ai_metadata=False,
        anonymize_reviewer_names=False,
        shape=shape,
    )


def _cell_values(out: ExportFile) -> set:
    wb = load_workbook(io.BytesIO(out.content))
    return {
        c.value
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for c in row
        if c.value is not None
    }


async def test_ni_appears_only_in_publication(db_session: AsyncSession) -> None:
    await _seed_finalized_absent_run(db_session)
    uid = str(SEED.primary_profile)
    complete = _cell_values(
        await build_export(db_session, _request(ExportShape.COMPLETE), user_id=uid)
    )
    pub = _cell_values(
        await build_export(db_session, _request(ExportShape.PUBLICATION), user_id=uid)
    )
    assert "NI" in pub and "NI = No information" in pub
    assert complete & {"NI", "NA", "NE", SCOPE_MARKER} == set()
