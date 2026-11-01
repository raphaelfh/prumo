"""Integration: the cell budget (spec §13, R9).

``estimate_sheet_cells`` must bound every rendered sheet from above, and
``build_export`` must refuse a build over ``EXPORT_CELL_BUDGET`` with the
stable ``EXPORT_TOO_LARGE`` code. Real Postgres, SEED graph; the SAVEPOINT
rollback drops every row written here.
"""

from __future__ import annotations

import io
from uuid import UUID, uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.extraction import ExtractionRun, ExtractionRunStage
from app.models.extraction_workflow import ExtractionReviewerDecisionType
from app.services.exports.extraction.run import (
    ExportRequest,
    ExportTooLargeError,
    build_export,
)
from app.services.exports.extraction.workbook import (
    ExportShape,
    allocate_titles,
    build_workbook,
    estimate_sheet_cells,
    matrix_fits,
)
from app.services.extraction_export_service import ExportLayout, ExportMode
from app.services.run_lifecycle_service import RunLifecycleService
from tests.integration.conftest import SEED, land_ai_proposal
from tests.integration.helpers.export_fixtures import export_service, make_run, record_decision

pytestmark = [pytest.mark.asyncio, pytest.mark.integration]

_RECORDS = 3


async def _extract_run(db: AsyncSession) -> ExtractionRun:
    run = await make_run(
        db,
        project_id=SEED.primary_project,
        article_id=SEED.primary_article,
        template_id=SEED.primary_template,
        profile_id=SEED.primary_profile,
    )
    if run.stage == ExtractionRunStage.PENDING.value:
        await RunLifecycleService(db).advance_stage(
            run_id=run.id, target_stage=ExtractionRunStage.EXTRACT, user_id=SEED.primary_profile
        )
    return run


async def _many_section_with_records(db: AsyncSession) -> tuple[UUID, list[UUID]]:
    """A root MANY section with one text field and three records on the seed article.

    The seed version snapshot is narrow, so the loader reads the live tables.
    """
    section_id, field_id = uuid4(), uuid4()
    await db.execute(
        text(
            "INSERT INTO public.extraction_entity_types "
            "(id, project_template_id, name, label, cardinality, "
            " parent_entity_type_id, sort_order, is_required, entry_label) "
            "VALUES (:id, :tid, :name, 'Arms', 'many', NULL, 1, false, 'arm')"
        ),
        {"id": str(section_id), "tid": str(SEED.primary_template), "name": "arms_budget"},
    )
    await db.execute(
        text(
            "INSERT INTO public.extraction_fields "
            "(id, entity_type_id, name, label, field_type, is_required) "
            "VALUES (:id, :etid, 'arm_name', 'Arm name', 'text', false)"
        ),
        {"id": str(field_id), "etid": str(section_id)},
    )
    instances: list[UUID] = []
    for i in range(_RECORDS):
        instance_id = uuid4()
        instances.append(instance_id)
        await db.execute(
            text(
                "INSERT INTO public.extraction_instances "
                "(id, project_id, template_id, entity_type_id, article_id, "
                " label, sort_order, created_by) "
                "VALUES (:id, :pid, :tid, :etid, :aid, :label, :so, :by)"
            ),
            {
                "id": str(instance_id),
                "pid": str(SEED.primary_project),
                "tid": str(SEED.primary_template),
                "etid": str(section_id),
                "aid": str(SEED.primary_article),
                "label": f"Arm {i + 1}",
                "so": i,
                "by": str(SEED.primary_profile),
            },
        )
    return field_id, instances


async def _two_reviewers_on_three_records(db: AsyncSession) -> ExportLayout:
    field_id, instances = await _many_section_with_records(db)
    run = await _extract_run(db)
    await land_ai_proposal(
        db, run_id=run.id, instance_id=instances[0], field_id=field_id,
        proposed_value={"value": "AI arm"},
    )  # fmt: skip
    for reviewer in (SEED.primary_profile, SEED.reviewer_profile):
        for i, instance_id in enumerate(instances):
            await record_decision(
                db, run_id=run.id, instance_id=instance_id, field_id=field_id,
                reviewer_id=reviewer, decision=ExtractionReviewerDecisionType.EDIT,
                value={"value": f"arm {i}"},
            )  # fmt: skip
    # all_users omits a run still in EXTRACT.
    await RunLifecycleService(db).advance_stage(
        run_id=run.id, target_stage=ExtractionRunStage.CONSENSUS, user_id=SEED.primary_profile
    )
    layout = await export_service(db, profile_id=SEED.primary_profile).resolve_layout(
        project_id=SEED.primary_project,
        template_id=SEED.primary_template,
        mode=ExportMode.ALL_USERS,
        article_ids=[SEED.primary_article],
        include_ai_metadata=True,
        anonymize_reviewer_names=False,
    )
    # Preconditions: two raters, three records, AI rows; else the bound is vacuous.
    assert len(layout.reviewers) == 2
    assert sum(1 for k in layout.value_map if k[2] == field_id) == 2 * _RECORDS
    assert layout.ai_proposal_rows
    return layout


async def test_estimate_cells_bounds_rendered_cells_per_sheet(db_session: AsyncSession) -> None:
    layout = await _two_reviewers_on_three_records(db_session)

    for shape in ExportShape:
        include_matrix = matrix_fits(layout)
        titles = allocate_titles(layout, shape, include_matrix=include_matrix)
        wb = load_workbook(io.BytesIO(build_workbook(layout, shape, include_matrix=include_matrix)))
        est = estimate_sheet_cells(layout, shape)
        sheet_of = {
            role: titles.fixed[role] if isinstance(role, str) else titles.levels[role]
            for role in est
        }

        assert set(sheet_of.values()) == set(wb.sheetnames), shape
        for role, n in est.items():
            ws = wb[sheet_of[role]]
            assert n >= ws.max_row * ws.max_column, (shape, role)


def _req(**kw: object) -> ExportRequest:
    base: dict[str, object] = {
        "project_id": SEED.primary_project,
        "template_id": SEED.primary_template,
        "mode": ExportMode.ALL_USERS,
        "article_ids": (SEED.primary_article,),
        "reviewer_id": None,
        "include_ai_metadata": False,
        "anonymize_reviewer_names": False,
        "shape": ExportShape.COMPLETE,
    }
    return ExportRequest(**{**base, **kw})  # type: ignore[arg-type]


async def test_oversized_build_raises_export_too_large(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    await _two_reviewers_on_three_records(db_session)
    monkeypatch.setattr("app.services.exports.extraction.run.EXPORT_CELL_BUDGET", 10)

    with pytest.raises(ExportTooLargeError) as exc:
        await build_export(db_session, _req(), user_id=str(SEED.primary_profile))

    assert exc.value.status_code == 422
    assert exc.value.code == "EXPORT_TOO_LARGE"
    assert exc.value.message.startswith("EXPORT_TOO_LARGE: ")


async def test_build_under_the_budget_succeeds(db_session: AsyncSession) -> None:
    await _two_reviewers_on_three_records(db_session)

    out = await build_export(db_session, _req(), user_id=str(SEED.primary_profile))

    assert "Data" in load_workbook(io.BytesIO(out.content)).sheetnames


def test_error_survives_celery_rebuild() -> None:
    e = ExportTooLargeError("EXPORT_TOO_LARGE: x")

    assert type(e)(*e.args).message == e.message
