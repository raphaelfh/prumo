"""HTTP-level tests for the extraction export endpoint family."""

from __future__ import annotations

import io
from unittest.mock import patch
from uuid import UUID, uuid4

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from tests.integration.conftest import SEED

_EP = "app.api.v1.endpoints.extraction_export"
_URL = "/api/v1/projects/{pid}/extraction-export"
_BODY = {
    "template_id": str(SEED.primary_template),
    "mode": "consensus",
    "article_ids": [],
    "shape": "dictionary",
}


async def test_dictionary_with_empty_article_ids_returns_200(
    db_client: AsyncClient, auth_as_profile: UUID
) -> None:
    assert auth_as_profile == SEED.primary_profile
    res = await db_client.post(_URL.format(pid=SEED.primary_project), json=_BODY)
    assert res.status_code == 200, res.text
    assert "Codebook" in load_workbook(io.BytesIO(res.content)).sheetnames


async def _insert_article(db: AsyncSession, project_id: UUID, *, title: str) -> UUID:
    article_id = uuid4()
    await db.execute(
        text(
            "INSERT INTO public.articles (id, project_id, title, row_version) "
            "VALUES (:id, :pid, :title, 1)"
        ),
        {"id": str(article_id), "pid": str(project_id), "title": title},
    )
    return article_id


async def test_an_origin_request_exposes_content_disposition(
    db_client: AsyncClient, auth_as_profile: UUID
) -> None:
    assert auth_as_profile == SEED.primary_profile
    res = await db_client.post(
        _URL.format(pid=SEED.primary_project),
        json=_BODY,
        headers={"Origin": "http://localhost:5173"},
    )
    assert res.status_code == 200, res.text
    assert "content-disposition" in res.headers["access-control-expose-headers"].lower()
    assert res.headers["content-disposition"].startswith('attachment; filename="')


async def test_foreign_article_is_404_on_an_async_request(
    db_client: AsyncClient, auth_as_profile: UUID, db_session: AsyncSession
) -> None:
    assert auth_as_profile == SEED.primary_profile
    foreign = await _insert_article(
        db_session, SEED.secondary_project, title="FOREIGN-SENTINEL-TITLE"
    )
    with (
        patch("app.worker.queue_health.is_queue_available", return_value=True),
        patch(f"{_EP}.export_extraction_task") as task,
    ):
        task.delay.return_value.id = "job-1"
        res = await db_client.post(
            _URL.format(pid=SEED.primary_project),
            json={
                **_BODY,
                "mode": "all_users",
                "shape": "complete",
                "article_ids": [str(foreign)],
            },
        )
    assert res.status_code == 404, res.text
    assert res.json()["error"]["code"] == "NOT_FOUND"
    assert "FOREIGN-SENTINEL-TITLE" not in res.text
    task.delay.assert_not_called()


async def test_foreign_template_is_404_on_an_async_request(
    db_client: AsyncClient, auth_as_profile: UUID
) -> None:
    assert auth_as_profile == SEED.primary_profile
    with (
        patch("app.worker.queue_health.is_queue_available", return_value=True),
        patch(f"{_EP}.export_extraction_task") as task,
    ):
        task.delay.return_value.id = "job-1"
        res = await db_client.post(
            _URL.format(pid=SEED.secondary_project), json={**_BODY, "mode": "all_users"}
        )
    assert res.status_code == 404, res.text
    assert res.json()["error"]["code"] == "NOT_FOUND"
    task.delay.assert_not_called()


@pytest.mark.parametrize("shape", ["complete", "publication"])
async def test_selected_only_with_empty_ids_stays_422_for_non_dictionary(
    db_client: AsyncClient, auth_as_profile: UUID, shape: str
) -> None:
    """R22: the deleted ``selected_only`` 400 was unreachable; 422 is unchanged."""
    assert auth_as_profile == SEED.primary_profile
    res = await db_client.post(
        _URL.format(pid=SEED.primary_project),
        json={**_BODY, "shape": shape, "article_scope": "selected_only"},
    )
    assert res.status_code == 422, res.text
